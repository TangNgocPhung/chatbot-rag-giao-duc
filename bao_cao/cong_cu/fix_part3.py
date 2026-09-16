# -*- coding: utf-8 -*-
"""Sửa nội dung, bố cục và hình thức tệp Word — phần 3.

Phần 3: chú thích bảng, hình minh hoạ Chương 3–4, mục đánh giá còn thiếu và
nội dung cho bốn phụ lục.
"""
import sys; sys.path.insert(0, '.')
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.table import Table

from docx_helpers import (paragraphs, find, find_contains, set_text, replace_text,
                          insert_after, insert_before, iter_blocks,
                          add_figure_after, add_table_caption)

doc = docx.Document('out_part2.docx')
FIG = 'fig/'
REPO = '/home/user/chatbot-rag-giao-duc/'
log = []
def L(m): log.append(m); print('·', m)
def body(ref, text): return insert_after(ref, doc, text)

# =========================================================================
# I. CHÚ THÍCH BẢNG (trước đây chỉ 1/12 bảng có chú thích nên danh mục bảng gần như trống)
# =========================================================================
def header_of(tbl):
    return ' | '.join(c.text.strip() for c in tbl.rows[0].cells)

CAPS = [
    ('Bạn là chuyên viên tra cứu chính sách giáo dục', 2, 1,
     'Mẫu prompt ràng buộc dùng cho bước sinh câu trả lời'),
    ('Thành phần | Thông số', 3, 1, 'Cấu hình phần cứng của máy thử nghiệm'),
    ('Phần mềm | Phiên bản | Vai trò', 3, 2, 'Phần mềm nền và vai trò của từng thành phần'),
    ('Đối tượng | Vị trí', 3, 3, 'Dữ liệu sinh ra khi vận hành và vị trí lưu trữ tương ứng'),
    ('__env1__', 3, 4, 'Nhóm tham số môi trường về mô hình và suy luận'),
    ('__env2__', 3, 5, 'Nhóm tham số môi trường về truy hồi và câu nhắc'),
    ('__env3__', 3, 6, 'Nhóm tham số môi trường về dữ liệu và chỉ mục'),
    ('__env4__', 3, 7, 'Nhóm tham số môi trường về máy chủ web và bảo mật'),
    ('Loại sự kiện | Nội dung', 3, 8, 'Các loại sự kiện NDJSON trả về trong một lượt hỏi – đáp'),
]
tables = [b for b in iter_blocks(doc) if isinstance(b, Table)]
env_tables = [t for t in tables if header_of(t).startswith('Biến | Mặc định')]
env_iter = iter(env_tables)
done = 0
for key, ch, num, text in CAPS:
    if key.startswith('__env'):
        tbl = next(env_iter)
    else:
        tbl = next((t for t in tables if key in header_of(t)), None)
    if tbl is None:
        L(f'CẢNH BÁO: không tìm thấy bảng cho chú thích "{text}"')
        continue
    add_table_caption(tbl, doc, text, ch, num)
    done += 1
L(f'Thêm chú thích tự đánh số cho {done} bảng (trước đây chỉ 1 trong 12 bảng có chú thích '
  f'nên "DANH MỤC BẢNG BIỂU" gần như trống)')

# =========================================================================
# J. CHƯƠNG 3 — hình minh hoạ
# =========================================================================
p = find_contains(doc, 'Hai môi trường này khác nhau về quy mô nhưng dùng chung một mã nguồn')
add_figure_after(p, doc, FIG + 'f10_trien_khai.png',
                 'Hai môi trường triển khai dùng chung một mã nguồn, chỉ khác tệp cấu hình môi trường',
                 chapter=3, number=1, width_cm=15.5)
L('Thêm hình so sánh hai môi trường triển khai vào mục 3.1.4')

p = find_contains(doc, 'Trạng thái này được giao diện hiển thị trực tiếp và dùng để khóa việc gửi câu hỏi')
add_figure_after(p, doc, FIG + 'f09_may_trang_thai.png',
                 'Máy trạng thái của dịch vụ RAG và ảnh hưởng tới trạng thái khung nhập câu hỏi',
                 chapter=3, number=2, width_cm=15.5)
L('Thêm hình máy trạng thái dịch vụ vào mục 3.2.2')

p = find_contains(doc, 'Trong lúc cập nhật, trạng thái chuyển sang updating')
add_figure_after(p, doc, FIG + 'f12_cap_nhat.png',
                 'Cơ chế cập nhật gia tăng chỉ mục dựa trên mã băm SHA-256 của từng tệp',
                 chapter=3, number=3, width_cm=15.5)
L('Thêm hình cơ chế cập nhật gia tăng vào Bước 5 của mục 3.4')

p = find_contains(doc, 'Thứ tự này có ý nghĩa với người dùng: danh sách nguồn xuất hiện trước')
add_figure_after(p, doc, FIG + 'f11_ndjson.png',
                 'Chuỗi sự kiện NDJSON của một lượt hỏi – đáp và ý nghĩa của thứ tự các sự kiện',
                 chapter=3, number=4, width_cm=15.5)
L('Thêm hình chuỗi sự kiện NDJSON vào Bước 6 của mục 3.4')

# =========================================================================
# K. CHƯƠNG 4 — mô tả kho tri thức, biểu đồ kết quả và mục hạn chế
# =========================================================================
anchor = find(doc, 'Tập dữ liệu và kịch bản kiểm thử', style='Heading 2')
new_h = insert_before(anchor, doc, 'Kho tri thức dùng trong thử nghiệm', style='Heading 2')
cur = new_h
for t in [
  'Toàn bộ kết quả trong chương này được đo trên kho tri thức thật của đề tài chứ không phải trên một '
  'tập chuẩn có sẵn. Tại thời điểm hoàn thiện, kho gồm 713 tệp tài liệu, chiếm khoảng 4,9 GB dung '
  'lượng gốc và được chuyển thành 25.761 vector trong chỉ mục FAISS.',

  'Về định dạng, PDF chiếm gần ba phần năm kho với 426 tệp (59,7%), tiếp đến là DOCX 164 tệp (23,0%), '
  'PPTX 72 tệp (10,1%), DOC 29 tệp (4,1%), XLSX 13 tệp (1,8%), MP4 5 tệp và XLS 4 tệp. Cơ cấu này '
  'giải thích vì sao khâu nạp dữ liệu phải có nhánh OCR cho PDF bản quét và nhánh phiên âm cho video: '
  'hai nhóm tệp đó tuy nhỏ về số lượng nhưng nếu bỏ qua thì một phần nội dung sẽ không bao giờ truy '
  'hồi được.',

  'Về nội dung, kho chia thành hai nhóm rõ rệt. Nhóm học liệu dạy học gồm 528 tệp (74,1%), trong đó '
  'sách giáo khoa và sách bài tập chiếm 302 tệp, kế hoạch bài dạy và bài giảng 151 tệp. Nhóm văn bản '
  'pháp quy và hành chính gồm 185 tệp (25,9%), trong đó Thông tư chiếm 108 tệp, Nghị định 28 tệp và '
  'Quyết định 24 tệp. Chính nhóm thứ hai tạo ra đặc thù khó nhất của bài toán: kho chứa đồng thời văn '
  'bản cũ và văn bản mới thay thế nó, hai bản gần như trùng nhau về ngữ nghĩa, nên riêng độ tương '
  'đồng vector không đủ để chọn đúng.',
]:
    cur = body(cur, t)
add_figure_after(cur, doc, FIG + 'repo_bieu_do_co_cau_kho.png',
                 'Cơ cấu 713 tài liệu trong kho tri thức theo định dạng tệp và theo thể loại nội dung',
                 chapter=4, number=1, width_cm=15.5)
L('Bổ sung mục 4.1 "Kho tri thức dùng trong thử nghiệm" kèm biểu đồ cơ cấu kho — bản cũ không mô tả '
  'quy mô và thành phần kho tài liệu')

p = find_contains(doc, 'Cách thiết kế này nhằm đánh giá khả năng kiểm soát phạm vi của hệ thống')
add_figure_after(p, doc, FIG + 'f16_bo_cau_hoi.png',
                 'Cơ cấu bộ 127 câu hỏi kiểm thử theo nhóm tài liệu nguồn và theo mục đích đánh giá',
                 chapter=4, number=2, width_cm=15.5)
L('Thêm biểu đồ cơ cấu bộ câu hỏi kiểm thử vào mục 4.2')

# Ba biểu đồ kết quả, đặt sau ảnh chụp kết quả chạy benchmark
p = find_contains(doc, 'Trên toàn bộ 97 câu hỏi, hệ thống đạt MRR@10 = 0,806')
f = add_figure_after(p, doc, FIG + 'f13_chi_so_theo_nhom.png',
                 'MRR@10, nDCG@10 và MAP@10 theo từng nhóm tài liệu nguồn',
                 chapter=4, number=4, width_cm=15.5)
f = add_figure_after(f, doc, FIG + 'f15_phan_bo_hang.png',
                 'Phân bố thứ hạng của nguồn đúng đầu tiên trên 97 câu hỏi có nhãn nguồn',
                 chapter=4, number=5, width_cm=14.5)
L('Thêm biểu đồ chỉ số theo nhóm và biểu đồ phân bố thứ hạng vào mục 4.4')

p = find_contains(doc, 'Nhóm video có kết quả thấp nhất')
add_figure_after(p, doc, FIG + 'f14_hit_at_k.png',
                 'Độ bao phủ Hit@K theo độ sâu danh sách kết quả ở từng nhóm tài liệu',
                 chapter=4, number=6, width_cm=15.5)
L('Thêm biểu đồ Hit@K theo nhóm tài liệu vào mục 4.4')

# --- Mục mới: hạn chế của phương pháp đánh giá ---------------------------
anchor = find(doc, 'KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN', style='Heading 1')
new_h = insert_before(anchor, doc, 'Hạn chế của phương pháp đánh giá', style='Heading 2')
cur = new_h
for t in [
  'Các con số ở mục trên cần được đọc kèm những giới hạn sau của chính cách đo.',

  'Thứ nhất, quy mô mẫu của một số nhóm quá nhỏ để khái quát hoá. Nhóm video chỉ có 5 câu hỏi và nhóm '
  'van_ban_doc chỉ có 6 câu, nên mỗi câu trả lời đúng hay sai đều làm chỉ số của nhóm dịch chuyển '
  'khoảng 17–20%. Khoảng tin cậy bootstrap 95% của MRR@10 toàn bộ là [0,735; 0,875] — đủ rộng để '
  'không nên so sánh hai nhóm nhỏ với nhau chỉ dựa trên thứ tự giá trị trung bình.',

  'Thứ hai, nhãn đánh giá là nhãn ở mức tệp chứ không phải ở mức đoạn. Trường nguon_mong_doi chỉ ghi '
  'một phần tên tệp nguồn, nên một kết quả được tính là đúng khi truy hồi trúng tệp, kể cả khi đoạn '
  'được lấy ra không chứa câu trả lời. Chỉ số IR vì vậy phản ánh khả năng tìm đúng tài liệu, không '
  'phản ánh trực tiếp chất lượng bằng chứng đưa vào prompt.',

  'Thứ ba, nhãn do chính nhóm tác giả gán nên có nguy cơ thiên lệch. Bộ câu hỏi được xây dựng bởi '
  'cùng nhóm đã thiết kế hệ thống, chưa qua đối chiếu chéo với người gán nhãn độc lập, và chưa đo độ '
  'đồng thuận giữa những người gán nhãn.',

  'Thứ tư, phép đo được thực hiện ở độ sâu 24 chunk trong khi hệ thống vận hành chỉ đưa 4 chunk vào '
  'prompt. Hai con số này trả lời hai câu hỏi khác nhau, nên Hit@10 = 90% không được hiểu là 90% câu '
  'hỏi có đủ bằng chứng đúng khi sinh câu trả lời.',

  'Thứ năm, đề tài chưa thực hiện nghiên cứu ablation. Vì tầng truy hồi luôn chạy nguyên khối hybrid '
  'kết hợp FAISS, BM25, RRF và xếp hạng lại heuristic, kết quả hiện tại không tách được phần đóng góp '
  'riêng của từng thành phần. Muốn kết luận rằng heuristic thực sự cải thiện thứ hạng thì phải dựng '
  'thêm các biến thể dense-only, keyword-only và hybrid-không-rerank rồi đo lại trên cùng bộ câu hỏi.',

  'Thứ sáu, chất lượng câu trả lời cuối cùng mới chỉ được kiểm soát bằng quy tắc. Tính hợp lệ của '
  'trích dẫn và mức độ có căn cứ của số liệu được đo bằng đối chiếu chuỗi, chưa có đánh giá của con '
  'người trên mẫu ngẫu nhiên và chưa dùng bộ đánh giá tự động theo hướng LLM-as-a-judge.',
]:
    cur = body(cur, t)
L('Bổ sung mục 4.5 "Hạn chế của phương pháp đánh giá" — phần bàn về độ tin cậy của kết quả mà bản cũ chưa có')

# --- Thống nhất dấu thập phân trong chương kết quả ----------------------
import re
in_ch4 = False
changed = 0
for b in iter_blocks(doc):
    if isinstance(b, Paragraph):
        if b.style.name == 'Heading 1':
            in_ch4 = b.text.strip().startswith('ĐÁNH GIÁ THỬ NGHIỆM')
        if not in_ch4:
            continue
        for r in b.runs:
            new = re.sub(r'(?<![\w.])0\.(\d{3})(?![\d.])', r'0,\1', r.text)
            if new != r.text:
                r.text = new; changed += 1
    else:
        for row in b.rows:
            for cell in row.cells:
                for p2 in cell.paragraphs:
                    for r in p2.runs:
                        new = re.sub(r'(?<![\w.])0\.(\d{3})(?![\d.])', r'0,\1', r.text)
                        if new != r.text:
                            r.text = new; changed += 1
L(f'Thống nhất dấu thập phân kiểu Việt Nam cho {changed} vị trí số liệu trong Chương 4 '
  f'(0.806 → 0,806), khớp với cách viết ở phần Kết luận')

# Bổ sung ghi chú kỹ thuật về nDCG dùng gain nhị phân
p = find_contains(doc, 'Trong benchmark, mức độ liên quan được xác định dựa trên sự phù hợp')
replace_text(p, 'Trong benchmark, mức độ liên quan được xác định dựa trên sự phù hợp của nguồn truy hồi với nguồn mong đợi.',
             'Trong benchmark của đề tài, mức độ liên quan là nhị phân (đúng nguồn mong đợi hoặc không), '
             'nên phần gain 2^rel − 1 rút gọn về 1 và công thức trở thành tổng của 1/log₂(hạng + 1) trên '
             'các vị trí trúng, chia cho giá trị lý tưởng tương ứng.')
L('Làm rõ nDCG trong bài dùng gain nhị phân (đúng với cài đặt trong chi_so_ir.py)')

doc.save('out_part3.docx')
print(f'\nĐÃ GHI out_part3.docx — {len(log)} thay đổi')
