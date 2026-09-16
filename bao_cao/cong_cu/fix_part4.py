# -*- coding: utf-8 -*-
"""Sửa nội dung, bố cục và hình thức tệp Word — phần 4: nội dung bốn phụ lục."""
import sys; sys.path.insert(0, '.')
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.table import Table

from docx_helpers import (paragraphs, find, find_contains, insert_after, insert_before,
                          iter_blocks, set_text)

doc = docx.Document('out_part3.docx')
REPO = '/home/user/chatbot-rag-giao-duc/'
log = []
def L(m): log.append(m); print('·', m)

def body(ref, text):
    return insert_after(ref, doc, text)

def code(ref, lines):
    """Chèn một khối mã nguồn giữ nguyên thụt đầu dòng."""
    p = insert_after(ref, doc, '', style='Source Code')
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        p.add_run(line.replace(' ', ' '))
    return p

def plain_caption(ref, text):
    p = insert_after(ref, doc, text, style='Caption')
    p.alignment = AL.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    return p

def picture(ref, path, width_cm):
    p = doc.add_paragraph()
    p.alignment = AL.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    ref._p.addnext(p._p)
    return p

# =========================================================================
# PHỤ LỤC A — mã nguồn chính
# =========================================================================
a = find(doc, 'Phụ lục A: Mã nguồn chính', style='Heading 2')
cur = body(a, 'Phụ lục này trích ba đoạn mã thể hiện rõ nhất các quyết định thiết kế đã trình bày '
              'ở Chương 2. Toàn bộ mã nguồn được công bố tại kho GitHub nêu ở Phụ lục C.')

cur = body(cur, 'A.1. Hợp nhất hai bảng xếp hạng bằng RRF (hybrid_retrieval.py)')
cur.runs[0].bold = True
cur = code(cur, [
 'RRF_K = 60                      # hằng số chuẩn trong công thức RRF',
 'SO_UNG_VIEN_MOI_RETRIEVER = 15  # top-k lấy từ mỗi nhánh trước khi hợp nhất',
 'SO_KET_QUA_CUOI = int(os.getenv("RAG_SO_BANG_CHUNG", "4"))',
 '',
 'def rrf_fusion(*danh_sach_ket_qua, k=RRF_K):',
 '    """Hợp nhất nhiều danh sách xếp hạng thành một, không phụ thuộc',
 '    thang điểm gốc của từng retriever."""',
 '    diem, tai_lieu, nguon_dong_gop = {}, {}, {}',
 '    for ten_retriever, ket_qua in danh_sach_ket_qua:',
 '        for hang, doc in enumerate(ket_qua):',
 '            khoa = _khoa_chunk(doc)',
 '            diem[khoa] = diem.get(khoa, 0.0) + 1.0 / (k + hang + 1)',
 '            tai_lieu.setdefault(khoa, doc)',
 '            nguon_dong_gop.setdefault(khoa, set()).add(ten_retriever)',
 '    xep_hang = sorted(diem.items(), key=lambda x: -x[1])',
 '    ...',
])
cur = body(cur, 'Khóa _khoa_chunk() được dựng theo nguồn và nội dung để nhận ra cùng một đoạn giữa hai '
                'nhánh, vì BM25 có thêm dòng tiêu đề nguồn vào văn bản lập chỉ mục.')

cur = body(cur, 'A.2. Định tuyến sang các công cụ tính trước khi vào RAG (rag_service.py)')
cur.runs[0].bold = True
cur = code(cur, [
 '# Thứ tự thử: công cụ có phạm vi hẹp nhất đi trước. tinh_toan nhận cả',
 '# những biểu thức trần trụi nên phải đứng cuối.',
 'for ten_cong_cu, cong_cu in (',
 '    ("tinh_luong", tinh_luong),',
 '    ("dinh_muc_tiet_day", dinh_muc_tiet_day),',
 '    ("danh_gia_hoc_sinh", danh_gia_hoc_sinh),',
 '    ("tinh_toan", tinh_toan),',
 '):',
 '    ket_qua_tinh = cong_cu.tra_loi(question)',
 '    if ket_qua_tinh is not None:',
 '        yield from self._tra_loi_cong_cu(',
 '            question, *ket_qua_tinh, ten_cong_cu',
 '        )',
 '        return',
])

cur = body(cur, 'A.3. Tham số sinh của mô hình ngôn ngữ (main.py)')
cur.runs[0].bold = True
cur = code(cur, [
 'llm_model = ten_model or os.getenv("RAG_LLM_MODEL", "qwen3.5:4b")',
 'return ChatOllama(',
 '    model=llm_model,',
 '    temperature=float(os.getenv("RAG_TEMPERATURE", "0.15")),',
 '    top_p=float(os.getenv("RAG_TOP_P", "0.9")),',
 '    repeat_penalty=float(os.getenv("RAG_REPEAT_PENALTY", "1.08")),',
 '    num_ctx=int(os.getenv("RAG_CONTEXT_LENGTH", "4096")),',
 '    num_predict=so_token_toi_da or int(os.getenv("RAG_MAX_OUTPUT_TOKENS", "420")),',
 '    reasoning=os.getenv("RAG_REASONING", "0") == "1",',
 '    keep_alive=os.getenv("RAG_KEEP_ALIVE", "2h"),',
 '    base_url=ollama_url,',
 ')',
])
cur = body(cur, 'Mọi tham số đều đọc từ biến môi trường và có giá trị mặc định an toàn ngay tại chỗ sử '
                'dụng, nên hệ thống chạy được ngay sau khi cài đặt mà không cần khai báo gì thêm.')
L('Điền nội dung cho Phụ lục A: ba trích đoạn mã nguồn chính (RRF, định tuyến công cụ tính, tham số sinh)')

# =========================================================================
# PHỤ LỤC B — ảnh chụp giao diện
# =========================================================================
b = find(doc, 'Phụ lục B: Ảnh chụp giao diện chạy thử', style='Heading 2')
cur = body(b, 'Các ảnh chụp dưới đây được lấy từ bản chạy thật của hệ thống. Ảnh giao diện chính, '
              'khung hội thoại kèm nguồn trích dẫn, thẻ trạng thái cập nhật chỉ mục và hộp thoại kho '
              'tài liệu đã được đặt trong Chương 2; phần này bổ sung hai màn hình còn lại.')
p = picture(cur, REPO + 'ui-desktop.png', 15.0)
cur = plain_caption(p, 'Hình B.1. Màn hình chào của ứng dụng trên máy tính để bàn: thẻ trạng thái hệ '
                       'thống, các chủ đề gợi ý và khung nhập câu hỏi')
p = picture(cur, REPO + 'ui-direct-file.png', 15.0)
cur = plain_caption(p, 'Hình B.2. Hỏi đáp trực tiếp trên một tệp vừa đính kèm: phạm vi tìm kiếm được '
                       'giới hạn trong chính tệp đó thay vì toàn bộ kho tri thức')
L('Điền nội dung cho Phụ lục B: hai ảnh chụp giao diện kèm chú thích (trước đây phụ lục để trống)')

# =========================================================================
# PHỤ LỤC C, D — thêm câu dẫn cho hai bảng liên kết
# =========================================================================
c = find(doc, 'Phụ lục C: Link github thực nghiệm', style='Heading 2')
body(c, 'Toàn bộ mã nguồn, bộ câu hỏi kiểm thử và công cụ đo chỉ số của đề tài được công bố tại kho '
        'GitHub dưới đây. Kho chỉ chứa mã nguồn; kho tài liệu giáo dục và chỉ mục FAISS được sinh ra '
        'tại máy chạy nên không đưa lên GitHub.')
d4 = find(doc, 'Phụ lục D: Link web thực nghiệm', style='Heading 2')
body(d4, 'Bản chạy thử phục vụ nhiều người dùng được triển khai trên máy chủ ảo theo mô tả ở mục 3.1.4. '
         'Mọi đường dẫn đều được bảo vệ bằng một lớp đăng nhập HTTP Basic, kể cả đầu cuối trạng thái.')
L('Thêm câu dẫn cho Phụ lục C và Phụ lục D (trước đây chỉ có bảng liên kết trơ trọi)')

doc.save('out_part4.docx')
print(f'\nĐÃ GHI out_part4.docx — {len(log)} thay đổi')
