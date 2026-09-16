# -*- coding: utf-8 -*-
"""Tiện ích thao tác trên tài liệu Word ở mức XML mà python-docx không cung cấp sẵn."""
import copy
import docx
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.table import Table

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def iter_blocks(doc):
    """Sinh lần lượt Paragraph và Table theo đúng thứ tự xuất hiện trong thân tài liệu."""
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, doc)
        elif child.tag == qn('w:tbl'):
            yield Table(child, doc)


def paragraphs(doc):
    return [b for b in iter_blocks(doc) if isinstance(b, Paragraph)]


def _n(s):
    """Chuẩn hoá khoảng trắng không ngắt để so khớp văn bản ổn định hơn."""
    return s.replace('\xa0', ' ').replace('\u2009', ' ')


def find(doc, prefix, style=None, nth=0):
    """Tìm đoạn văn có nội dung bắt đầu bằng ``prefix``. Báo lỗi nếu không thấy."""
    hits = [p for p in paragraphs(doc)
            if _n(p.text).strip().startswith(_n(prefix)) and (style is None or p.style.name == style)]
    if not hits:
        raise LookupError(f'Không tìm thấy đoạn bắt đầu bằng: {prefix!r}')
    return hits[nth]


def find_contains(doc, needle, nth=0):
    hits = [p for p in paragraphs(doc) if _n(needle) in _n(p.text)]
    if not hits:
        raise LookupError(f'Không tìm thấy đoạn chứa: {needle!r}')
    return hits[nth]


def set_text(par, text):
    """Thay toàn bộ nội dung đoạn, giữ định dạng của run đầu tiên."""
    runs = par.runs
    if not runs:
        par.add_run(text)
        return par
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return par


def replace_text(par, old, new):
    """Thay chuỗi con trong đoạn, kể cả khi chuỗi bị cắt ngang qua nhiều run.

    Chỉ chạm vào đúng các run chứa đoạn khớp nên các phần tử khác của đoạn
    (trường trích dẫn, công thức, siêu liên kết) được giữ nguyên.
    """
    runs = par.runs
    if not runs:
        return False
    texts = [_n(r.text) for r in runs]
    full = ''.join(texts)
    pos = full.find(_n(old))
    if pos < 0:
        return False
    end = pos + len(_n(old))
    start_i = start_off = None
    acc = 0
    for i, t in enumerate(texts):
        if start_i is None and acc + len(t) > pos:
            start_i, start_off = i, pos - acc
        if acc + len(t) >= end:
            end_i, end_off = i, end - acc
            break
        acc += len(t)
    else:
        return False
    head = texts[start_i][:start_off]
    tail = texts[end_i][end_off:]
    runs[start_i].text = head + new + (tail if start_i == end_i else '')
    for i in range(start_i + 1, end_i + 1):
        runs[i].text = tail if i == end_i else ''
    return True


def _new_p(doc, text, style=None, align=None, indent_first=None, bold=False, italic=False):
    p = doc.add_paragraph()                       # tạm thêm vào cuối, sẽ dời sau
    if style:
        p.style = doc.styles[style]
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
    if align is not None:
        p.alignment = align
    if indent_first is not None:
        p.paragraph_format.first_line_indent = indent_first
    return p


def insert_after(ref, doc, text, style=None, align=None, indent_first=None,
                 bold=False, italic=False):
    """Chèn một đoạn mới ngay sau ``ref`` (Paragraph hoặc Table)."""
    p = _new_p(doc, text, style, align, indent_first, bold, italic)
    anchor = ref._p if isinstance(ref, Paragraph) else ref._tbl
    anchor.addnext(p._p)
    return p


def insert_before(ref, doc, text, style=None, align=None, indent_first=None,
                  bold=False, italic=False):
    p = _new_p(doc, text, style, align, indent_first, bold, italic)
    anchor = ref._p if isinstance(ref, Paragraph) else ref._tbl
    anchor.addprevious(p._p)
    return p


def insert_paragraphs_after(ref, doc, items):
    """items: list các tuple (text, style) hoặc chuỗi. Trả về đoạn cuối cùng đã chèn."""
    cur = ref
    for it in items:
        if isinstance(it, tuple):
            text, style = it[0], it[1]
            kw = it[2] if len(it) > 2 else {}
        else:
            text, style, kw = it, None, {}
        cur = insert_after(cur, doc, text, style, **kw)
    return cur


CAPTION_XML = """<w:p {ns}>
 <w:pPr><w:pStyle w:val="Caption"/>{keep}<w:jc w:val="center"/><w:ind w:firstLine="0"/></w:pPr>
 <w:r><w:t xml:space="preserve">{kind} </w:t></w:r>
 <w:r><w:fldChar w:fldCharType="begin"/></w:r>
 <w:r><w:instrText xml:space="preserve"> STYLEREF 1 \\s </w:instrText></w:r>
 <w:r><w:fldChar w:fldCharType="separate"/></w:r>
 <w:r><w:rPr><w:noProof/></w:rPr><w:t>{ch}</w:t></w:r>
 <w:r><w:fldChar w:fldCharType="end"/></w:r>
 <w:r><w:t>.</w:t></w:r>
 <w:r><w:fldChar w:fldCharType="begin"/></w:r>
 <w:r><w:instrText xml:space="preserve"> SEQ {kind} \\* ARABIC \\s 1 </w:instrText></w:r>
 <w:r><w:fldChar w:fldCharType="separate"/></w:r>
 <w:r><w:rPr><w:noProof/></w:rPr><w:t>{num}</w:t></w:r>
 <w:r><w:fldChar w:fldCharType="end"/></w:r>
 <w:r><w:t xml:space="preserve">. {text}</w:t></w:r>
</w:p>"""


def make_caption(kind, chapter, number, text, keep_next=False):
    """Tạo đoạn chú thích dùng trường STYLEREF + SEQ để Word tự đánh số lại."""
    xml = CAPTION_XML.format(ns=nsdecls('w'), kind=kind, ch=chapter, num=number,
                             text=_esc(text), keep='<w:keepNext/>' if keep_next else '')
    return parse_xml(xml)


def _esc(s):
    return (s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))


def add_figure_after(ref, doc, image_path, caption_text, chapter, number,
                     width_cm=15.5, kind='Hình'):
    """Chèn ảnh (căn giữa) kèm chú thích tự đánh số ngay sau ``ref``."""
    p = doc.add_paragraph()
    p.alignment = 1                                   # CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(image_path, width=Cm(width_cm))
    anchor = ref._p if isinstance(ref, Paragraph) else ref._tbl
    anchor.addnext(p._p)
    cap = make_caption(kind, chapter, number, caption_text)
    p._p.addnext(cap)
    return Paragraph(cap, doc)


def add_table_caption(tbl, doc, caption_text, chapter, number):
    cap = make_caption('Bảng', chapter, number, caption_text, keep_next=True)
    tbl._tbl.addprevious(cap)
    return Paragraph(cap, doc)


def center_image_paragraphs(doc):
    """Căn giữa và bỏ thụt dòng đầu ở mọi đoạn chỉ chứa ảnh."""
    n = 0
    for p in paragraphs(doc):
        if p._p.findall('.//' + qn('a:blip')) and not p.text.strip():
            p.alignment = 1
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            n += 1
    return n


def set_page_margins(doc, top=1418, bottom=1418, left=1701, right=1134):
    n = 0
    for sect in doc.element.body.iter(qn('w:sectPr')):
        mar = sect.find(qn('w:pgMar'))
        if mar is None:
            continue
        mar.set(qn('w:top'), str(top))
        mar.set(qn('w:bottom'), str(bottom))
        mar.set(qn('w:left'), str(left))
        mar.set(qn('w:right'), str(right))
        n += 1
    return n


def enable_update_fields(doc):
    """Bật cờ để Word cập nhật lại toàn bộ trường (mục lục, số hình, số bảng) khi mở tệp."""
    settings = doc.settings.element
    for tag in settings.findall(qn('w:updateFields')):
        settings.remove(tag)
    el = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
    settings.append(el)
    # đánh dấu các trường TOC là "bẩn" để chắc chắn được tính lại
    for instr in doc.element.body.iter(qn('w:instrText')):
        if instr.text and instr.text.strip().startswith(('TOC', 'PAGEREF')):
            r = instr.getparent()
            prev = r.getprevious()
            if prev is not None:
                fld = prev.find(qn('w:fldChar'))
                if fld is not None and fld.get(qn('w:fldCharType')) == 'begin':
                    fld.set(qn('w:dirty'), 'true')


def toc_field(doc, instr, placeholder):
    """Tạo đoạn chứa một trường TOC (ví dụ danh mục hình)."""
    xml = f"""<w:p {nsdecls('w')}>
     <w:pPr><w:pStyle w:val="TableofFigures"/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9072"/></w:tabs><w:ind w:firstLine="0"/><w:rPr><w:noProof/></w:rPr></w:pPr>
     <w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>
     <w:r><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>
     <w:r><w:fldChar w:fldCharType="separate"/></w:r>
     <w:r><w:t>{_esc(placeholder)}</w:t></w:r>
     <w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>"""
    return parse_xml(xml)
